"""Dataset loaders for the parameter-search harness. Every loader returns a ``fastbt.Market``.

Data lives in ``data/param_search/`` (git-ignored; archived on titanserver under
``/home/TitanTrade/archive/param_search/``). Treat the files as untrusted input — run with ``python -I``.

| key | universe | benchmark | span | hindsight |
|---|---|---|---|---|
| A_alpaca_names_2016   | the 15 live names (Alpaca SIP)            | SPY            | 2016 → | names chosen in 2026 |
| D_alpaca_sectors_2016 | 11 SPDR sector ETFs (Alpaca SIP)          | SPY            | 2016 → | none |
| B_yahoo_names_1999    | the 15 live names (Yahoo adj. close)      | SPY            | 1999 → | names chosen in 2026 |
| C_yahoo_sectors_1999  | 11 SPDR sector ETFs (Yahoo adj. close)    | SPY            | 1999 → | none |
| E_sp500_2016          | S&P 500 members as of 2016-01-04 (502 with Alpaca bars, delisted names included) | SPY | 2016 → | none (point-in-time list) |
| F_french_49ind_1926   | Kenneth French 49 industry portfolios, value-weighted daily total returns → index levels | French market factor (Mkt-RF + RF) | 1926 → | none |

Not a ``load`` key: ``alpaca(sym, "intl")`` reads the 39 region / country ETFs (roadmap E1, 2016 →, dividend-adjusted) and
``french_regional(region, kind)`` Kenneth French's regional daily factors (``external/french_intl/``, USD, 1990 →);
``cboe(name)`` a Cboe index history from the external library (``external/cboe/``, option-strategy benchmarks 1986 →, VIX family);
``french_levels(file, sub="external/french")`` the Decision 072 French sorts (profitability, investment, size × profitability, 10 / 30
industries), ``french_factors(file, sub)`` any French factor file (3 / 5 factors, momentum) as decimal returns, ``aqr_qmj(column)``
AQR's Quality-Minus-Junk daily factor by country / region (openpyxl).
"""
import datetime
import json
import os

import fastbt

S = os.environ.get("PARAM_SEARCH_DATA", "data/param_search")
NAMES = ["CRWD", "ANET", "LLY", "DXCM", "HCA", "JPM", "GS", "DVN", "WMT", "URI", "GE", "DASH", "DECK", "FCX", "EQIX"]
SECT = ["XLK", "XLV", "XLF", "XLE", "XLI", "XLY", "XLB", "XLRE", "XLP", "XLU", "XLC"]


def alpaca(sym, sub="historical_long"):
    return {b["date"]: float(b["close"]) for b in json.load(open(f"{S}/alpaca_hist/{sub}/{sym}.json"))}


def yahoo(sym):
    r = json.load(open(f"{S}/yahoo/{sym}.json"))["chart"]["result"][0]
    out = {}
    for t, c in zip(r["timestamp"], r["indicators"]["adjclose"][0]["adjclose"]):
        if c is not None and c > 0:
            out[datetime.datetime.fromtimestamp(t, datetime.timezone.utc).strftime("%Y-%m-%d")] = float(c)
    return out


def _french_section(path, header_prefix):
    """Rows of the first section whose title starts with ``header_prefix``; dates YYYYMMDD → YYYY-MM-DD; -99.99/-999 → None."""
    lines = open(path, encoding="latin-1").read().splitlines()
    i = next(k for k, line in enumerate(lines) if line.strip().startswith(header_prefix))
    h = i if lines[i].startswith(",") else i + 1          # the factors file has no title line above its header
    cols = [c.strip() for c in lines[h].split(",")][1:]
    rows = {}
    for line in lines[h + 1:]:
        if not line[:1].isdigit():
            break
        parts = [p.strip() for p in line.split(",")]
        d = f"{parts[0][:4]}-{parts[0][4:6]}-{parts[0][6:8]}"
        rows[d] = [None if v in ("-99.99", "-999") else float(v) for v in parts[1:]]
    return cols, rows


def french_market():
    """Market total-return index from the daily factors file (Mkt-RF + RF, in %)."""
    cols, rows = _french_section(f"{S}/french/F-F_Research_Data_Factors_daily.csv", ",Mkt-RF")
    im, ir = cols.index("Mkt-RF"), cols.index("RF")
    level, out = 100.0, {}
    for d in sorted(rows):
        level *= 1 + (rows[d][im] + rows[d][ir]) / 100
        out[d] = level
    return out


def french_factors(filename, sub="french"):
    """A French factor file (header starts with ",") → {column: {date: decimal return}}; -99.99 rows are skipped."""
    cols, rows = _french_section(f"{S}/{sub}/{filename}", ",")
    return {c: {d: rows[d][j] / 100 for d in sorted(rows) if rows[d][j] is not None} for j, c in enumerate(cols)}


def aqr_qmj(column="USA"):
    """AQR's Quality-Minus-Junk daily factor (``external/aqr/Quality-Minus-Junk-Factors-Daily.xlsx``, sheet "QMJ Factors",
    one column per country / region, decimal returns) → {date: return}. Needs openpyxl (``uv run --with openpyxl``)."""
    import openpyxl
    ws = openpyxl.load_workbook(f"{S}/external/aqr/Quality-Minus-Junk-Factors-Daily.xlsx", read_only=True)["QMJ Factors"]
    out, j = {}, None
    for row in ws.iter_rows(values_only=True):
        if j is None:
            if row and row[0] == "DATE":
                j = list(row).index(column)
            continue
        if not row or row[0] is None or row[j] is None:
            continue
        d = row[0] if isinstance(row[0], str) else row[0].strftime("%m/%d/%Y")
        m, dd, y = d[:10].split("/")
        out[f"{y}-{m}-{dd}"] = float(row[j])
    return out


def aqr_monthly(filename, sheet, header_test):
    """Every numeric column of an AQR monthly sheet → {column: {"YYYY-MM": decimal return}}. ``header_test(row)`` picks
    the header row (the Century file's starts with "Date", the TSMOM file's with None then "TSMOM"); a cell that is not
    a number is skipped. Needs openpyxl (``uv run --with openpyxl``)."""
    import openpyxl
    ws = openpyxl.load_workbook(f"{S}/external/aqr/{filename}", read_only=True)[sheet]
    cols, out = None, {}
    for row in ws.iter_rows(values_only=True):
        if cols is None:
            if row and header_test(row):
                cols = [(j, str(c)) for j, c in enumerate(row) if j > 0 and c not in (None, "")]
                out = {c: {} for _, c in cols}
            continue
        if not row or row[0] in (None, ""):
            continue
        d = row[0] if isinstance(row[0], str) else row[0].strftime("%Y-%m-%d")
        ym = f"{d[6:10]}-{d[0:2]}" if "/" in d else d[:7]
        for j, c in cols:
            v = row[j] if j < len(row) else None
            if isinstance(v, (int, float)):
                out[c][ym] = float(v)
    return out


def french_levels(filename, sub="french"):
    """Value-weighted daily total-return index per column of a French portfolio file (first VW section)."""
    cols, rows = _french_section(f"{S}/{sub}/{filename}", "Average Value Weighted Returns")
    levels = {c: 100.0 for c in cols}; out = {c: {} for c in cols}
    for d in sorted(rows):
        for j, c in enumerate(cols):
            r = rows[d][j]
            if r is None:
                continue
            levels[c] *= 1 + r / 100
            out[c][d] = levels[c]
    return out


def french_industries():
    """Value-weighted daily total-return index per industry; a missing return leaves that day out of the series."""
    cols, rows = _french_section(f"{S}/french/49_Industry_Portfolios_Daily.csv", "Average Value Weighted Returns")
    levels = {c: 100.0 for c in cols}; out = {c: {} for c in cols}
    for d in sorted(rows):
        for j, c in enumerate(cols):
            r = rows[d][j]
            if r is None:
                continue
            levels[c] *= 1 + r / 100
            out[c][d] = levels[c]
    return out


def french_regional(region, kind):
    """Daily decimal returns from French's regional files: ``kind`` "mom" → WML, "mkt" → Mkt-RF + RF;
    ``region`` = Developed_ex_US / Europe / Japan / Asia_Pacific_ex_Japan / Developed (the zips unpacked in place)."""
    if kind == "mom":
        cols, rows = _french_section(f"{S}/external/french_intl/{region}_MOM_Factor_Daily.csv", ",WML")
        i = cols.index("WML")
        return {d: rows[d][i] / 100 for d in sorted(rows) if rows[d][i] is not None}
    cols, rows = _french_section(f"{S}/external/french_intl/{region}_3_Factors_Daily.csv", ",Mkt-RF")
    im, ir = cols.index("Mkt-RF"), cols.index("RF")
    return {d: (rows[d][im] + rows[d][ir]) / 100 for d in sorted(rows) if rows[d][im] is not None}


def cboe(name):
    """Cboe daily index history (``external/cboe/<NAME>_History.csv``, MM/DD/YYYY) → {date: close}; the VIX-family files
    carry OHLC (the last column is the close), the strategy indices one column."""
    out = {}
    for line in open(f"{S}/external/cboe/{name}_History.csv").read().splitlines()[1:]:
        p = line.strip().split(",")
        if len(p) < 2 or not p[0][:1].isdigit():
            continue
        m, d, y = p[0].split("/")
        out[f"{y}-{m}-{d}"] = float(p[-1])
    return out


def sp500_2016_symbols():
    return json.load(open(f"{S}/universes/sp500_2016-01-04.json"))


def load(key):
    if key == "A_alpaca_names_2016":
        return fastbt.Market(alpaca("SPY"), {n: alpaca(n) for n in NAMES}, {s: alpaca(s) for s in SECT})
    if key == "D_alpaca_sectors_2016":
        return fastbt.Market(alpaca("SPY"), {s: alpaca(s) for s in SECT}, {s: alpaca(s) for s in SECT})
    if key == "B_yahoo_names_1999":
        return fastbt.Market(yahoo("SPY"), {n: yahoo(n) for n in NAMES}, {s: yahoo(s) for s in SECT})
    if key == "C_yahoo_sectors_1999":
        return fastbt.Market(yahoo("SPY"), {s: yahoo(s) for s in SECT}, {s: yahoo(s) for s in SECT})
    if key == "E_sp500_2016":
        syms = [s for s in sp500_2016_symbols() if os.path.exists(f"{S}/alpaca_hist/sp500_2016/{s}.json")]
        return fastbt.Market(alpaca("SPY"), {s: alpaca(s, "sp500_2016") for s in syms}, {s: alpaca(s) for s in SECT})
    if key == "F_french_49ind_1926":
        ind = french_industries()
        return fastbt.Market(french_market(), ind, ind)
    raise KeyError(key)


START = {"A_alpaca_names_2016": "2016-01-04", "D_alpaca_sectors_2016": "2016-01-04", "B_yahoo_names_1999": "1999-01-04",
         "C_yahoo_sectors_1999": "1999-01-04", "E_sp500_2016": "2016-01-04", "F_french_49ind_1926": "1927-07-01"}
SPLIT = {"A_alpaca_names_2016": "2022-01-01", "D_alpaca_sectors_2016": "2022-01-01", "B_yahoo_names_1999": "2018-01-01",
         "C_yahoo_sectors_1999": "2018-01-01", "E_sp500_2016": "2022-01-01", "F_french_49ind_1926": "1990-01-01"}
